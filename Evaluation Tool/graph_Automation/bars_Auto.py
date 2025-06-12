import openpyxl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import FuncFormatter
import os

# ==========================
# SHEET AND FILE SELECTION
# ==========================

sheet_name = "3_SoA_Processes"

file_name = "performance_Analysis.xlsx"
#file_name = "cost_Analysis.xlsx"

# ==========================
# FORMATTER FUNCTION
# ==========================

def no_scientific(x, pos):
    """Custom formatter for Y-axis to avoid scientific notation and show commas."""
    return f'{x:,.0f}'

# ==========================
# LOAD WORKBOOK AND SHEET
# ==========================

workbook = openpyxl.load_workbook(file_name, data_only=True)
sheet = workbook[sheet_name]

# ==========================
# DATA COLLECTION
# ==========================

x_labels = []  # X-axis labels for the plot
res = []       # Collected data results

if file_name == "performance_Analysis.xlsx":

    # ======================
    # CASE: performance_Analysis
    # ======================

    for table_name in sheet.tables:
        table = sheet.tables[table_name]

        # Extract header indexes
        header_values = [cell.value for cell in sheet[table.ref][0]]
        avg_index = header_values.index("Average")
        functions_index = header_values.index("Function/Endpoint")
        type_index = header_values.index("Type")

        # Iterate through table rows
        summo = 0
        res2 = []
        done = False

        for row in sheet[table.ref][1:]:
            if row[type_index].value == "Rest":
                summo += row[avg_index].value

                if row[functions_index].value == "saveModel":
                    res2.append(int(summo))
                    summo = 0

                if row[functions_index].value == "attributesCertification":
                    res2.append(int(summo))
                    summo = 0

                elif "decrypt" in row[functions_index].value and not done:
                    res2.append(int(summo - row[avg_index].value))
                    done = True
                    summo = row[avg_index].value

        res2.append(int(summo))  # Append the remaining value
        res.append(res2)


else:
    # ======================
    # CASE: cost_Analysis
    # ======================

    for table_name in sheet.tables:
        table = sheet.tables[table_name]

        # Extract header indexes
        header_values = [cell.value for cell in sheet[table.ref][0]]
        gas_index = header_values.index("Gas")
        functions_index = header_values.index("Function")

        summo = 0
        res2 = []
        done = False
        done2 = False

        # Iterate through table rows
        for row in sheet[table.ref][1:]:
            func = row[functions_index].value

            if func == "setIPFSLink" and not done:
                res2.append(summo)
                done = True
                summo = 0

            elif func == "notifyAuthorities" and not done2:
                res2.append(summo)
                done2 = True
                summo = 0

            summo += row[gas_index].value

        res2.append(summo)  # Append the remaining gas value
        res.append(res2)

# ==========================
# DATA VISUALIZATION - STACKED BAR CHART
# ==========================

data = np.array(res)
num_points = data.shape[0]

# ==========================
# MODIFIED X-POSITION CALCULATION
# ==========================
bar_width = 0.5
delta = 0.2

spacing = 1.2  # Increase this number for more space between bars
x = np.arange(num_points) * spacing

# Create the plot
fig, ax = plt.subplots(figsize=(10, 6))

# Define color palette
fixed_colors = ['#ef476f', '#ffd166', '#06d6a0', '#118ab2', '#073b4c']
alpha_val = 0.7

# Set phase labels and colors based on the file type
if file_name == "performance_Analysis.xlsx":
    phase_labels = ['Configure', 'Instantiate', 'Transact', 'Inspect']
    colors = fixed_colors[:4]
else:
    phase_labels = ['Instantiate', 'Transact', 'Inspect']
    colors = fixed_colors[1:4]

# Initialize the bottom for stacking
bottom = np.zeros(num_points)

# Plot each phase as a stacked bar with adjusted positions
for phase in range(data.shape[1]):
    ax.bar(x, data[:, phase], bottom=bottom, color=colors[phase],
           alpha=alpha_val, label=phase_labels[phase],
           edgecolor='grey', linewidth=1, width=bar_width)
    bottom += data[:, phase]

# ==========================
# AXIS CONFIGURATION
# ==========================
ax.set_xticks(x)
x_labels = ["X-ray", "Retail", "Incident"]
ax.set_xticklabels(x_labels)

# Set axis limits with specified edge padding
ax.set_xlim(left=delta - bar_width - 1.9, 
           right=(num_points - 1) + delta + bar_width * 2)

xlabel = 'Process'
file_label = '3_SoA_Processes'
ax.set_xlabel(xlabel, fontsize=24)

# Y-axis configuration (unchanged)
if file_name == "performance_Analysis.xlsx":
    ax.set_ylabel("Cumulative exec. time [ms]", fontsize=24)
    tick_step = 2_000
    pdf_prefix = "time"
else:
    ax.set_ylabel("Cumulative gas used [GU]", fontsize=24)
    tick_step = 1_000_000
    pdf_prefix = "gas"

total = data.sum(axis=1)
y_min = 0
y_max = total.max()

y_min_adjusted = (y_min // tick_step) * tick_step
y_max_adjusted = ((y_max // tick_step) + 1) * tick_step

y_ticks = np.arange(y_min_adjusted, y_max_adjusted + tick_step, tick_step)
ax.set_yticks(y_ticks)
ax.set_ylim(bottom=y_min_adjusted, top=y_max_adjusted)
ax.yaxis.set_major_formatter(FuncFormatter(no_scientific))

ax.grid(axis='y', linestyle='--', alpha=0.7)
ax.tick_params(axis='x', labelsize=18)
ax.tick_params(axis='y', labelsize=18)
ax.legend(title="Functionalities", loc='upper left', fontsize=16, title_fontsize=18)

plt.tight_layout()

# ==========================
# SAVE AS PDF
# ==========================
pdf_filename = f"{pdf_prefix}_{file_label}.pdf"
output_folder = "pdf_Outputs"
os.makedirs(output_folder, exist_ok=True)
pdf_path = os.path.join(output_folder, pdf_filename)
plt.savefig(pdf_path, format='pdf')
print(f"Figure saved as: {pdf_path}")

plt.show()
