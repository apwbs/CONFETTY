import openpyxl
import matplotlib.pyplot as plt
import itertools
import numpy as np
from matplotlib.ticker import FuncFormatter
import random

sheet_name = "base_Case"
sheet_name = "Encryptors"
sheet_name = "message_Size"
sheet_name = "Looping"
sheet_name = "parallel_Opening"
sheet_name = "parallel_Opening_Closing"
sheet_name = "exclusive_Opening"
sheet_name = "exclusive_Opening_Closing"


file_name = "performance_Analysis.xlsx"
file_name = "cost_Analysis.xlsx"

workbook = openpyxl.load_workbook(file_name, data_only=True)
sheet = workbook[sheet_name]

if file_name == "performance_Analysis.xlsx":
    if sheet_name == "base_Case":
    
        for table_name in sheet.tables:
            functions = []
            average = []
            table = sheet.tables[table_name]
                
            header_values = [cell.value for cell in sheet[table.ref][0]]
            avg_index = header_values.index("Average")
            type_index = header_values.index("Type")
            function_index = header_values.index("Function/Endpoint")
            for row in sheet[table.ref][1:]:
                if row[type_index].value == "Rest":
                    average.append(row[avg_index].value)
                    functions.append(row[function_index].value)
        res = [int(x) for x in itertools.accumulate(average)]
        phases = []
        val = 0
        done = False
        for e in range(len(average)):
            val += average[e]
            if e == 0:
                phases.append(int(val))
            if functions[e] == "attributesCertification":
                phases.append(int(val))
            if "decrypt" in functions[e] and done == False:
                done = True
                phases.append(int(val-average[e]))
        phases.append(int(val))
        x_labels = ['Configure', 'Instantiate', 'Transact', 'Inspect']
        
        
        plt.plot(res, color=(random.random(), random.random(), random.random()), label='Healthcare process')
        xtick_positions = [res.index(phase) for phase in phases]
        plt.xticks(xtick_positions, x_labels)
        plt.xlabel('Functionality')
        for phase in phases:
            plt.axvline(x=res.index(phase), color='black', linestyle=':', linewidth=1)
        
    elif sheet_name == "Encryptors" or sheet_name == "Looping":
        res = []
        x_labels = []
        for table_name in sheet.tables:
            summo = 0
            table = sheet.tables[table_name]
            if sheet_name == "Encryptors":
                x_labels.append(table_name.split("_")[0])
                plt.xlabel('Number of encryptors')
            elif sheet_name == "Looping":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Loops')
            header_values = [cell.value for cell in sheet[table.ref][0]]
            avg_index = header_values.index("Average")
            type_index = header_values.index("Type")
            for row in sheet[table.ref][1:]:
                if row[type_index].value == "Rest":
                    summo += row[avg_index].value
            res.append(int(summo))
        
        plt.plot(x_labels, res, color=(random.random(), random.random(), random.random()), label='Healthcare process')
        
    else:
        res = []
        x_labels = []
        for table_name in sheet.tables:
            summo = 0
            row_new = 2
            table = sheet.tables[table_name]
            if sheet_name == "message_Size":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Size duplication')
            elif sheet_name == "parallel_Opening":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Openings')
            elif sheet_name == "exclusive_Opening":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Openings')
            elif sheet_name == "parallel_Opening_Closing":
                x_labels.append(table_name.split("_")[2])
                plt.xlabel('Openings and Closings')
            elif sheet_name == "exclusive_Opening_Closing":
                x_labels.append(table_name.split("_")[2])
                plt.xlabel('Openings and Closings')
            header_values = [cell.value for cell in sheet[table.ref][0]]
            avg_index = header_values.index("Average")
            for row in sheet[table.ref][1:]:
                type_index = sheet.cell(row = row_new, column = 1).value
                if type_index == "Rest":
                    summo += row[avg_index].value
                row_new += 1
            res.append(int(summo))
        plt.plot(x_labels, res, color=(random.random(), random.random(), random.random()), label='Healthcare process')
        
    y_ticks = np.linspace(min(res), max(res), 7)
    plt.yticks(y_ticks)
    plt.gca().yaxis.set_major_formatter(FuncFormatter(lambda x, _: f'{int(x)}'))
    min_val = min(res)
    plt.fill_between(range(len(res)), min_val, res, color=(random.random(), random.random(), random.random()), alpha=0.3) 
    plt.grid(axis='y', linestyle='-', alpha=0.5)
    plt.ylabel('Cumulative exec. time [ms]')
    plt.legend()
    plt.show()

else:
    if sheet_name == "base_Case":
        for table_name in sheet.tables:
            functions = []
            gasses = []
            table = sheet.tables[table_name]
                
            header_values = [cell.value for cell in sheet[table.ref][0]]
            gas_index = header_values.index("Gas")
            function_index = header_values.index("Function")
            for row in sheet[table.ref][1:]:
                gasses.append(row[gas_index].value)
                functions.append(row[function_index].value)
        res = [int(x) for x in itertools.accumulate(gasses)]
        phases = []
        val = 0
        done = False
        done2 = False
        for e in range(len(res)):
            val += gasses[e]
            if functions[e] == "setIPFSLink" and done == False:
                phases.append(int(val) - gasses[e])
                done = True
            if functions[e] == "notifyAuthorities" and done2 == False:
                phases.append(int(val - gasses[e]))
                done2 = True
        phases.append(int(val))
        x_labels = ['Instantiate', 'Transact', 'Inspect']

        plt.plot(res, color=(random.random(), random.random(), random.random()), label='Healthcare process')
        xtick_positions = [res.index(phase) for phase in phases]
        plt.xticks(xtick_positions, x_labels)
        plt.xlabel('Functionality')
        for phase in phases:
            plt.axvline(x=res.index(phase), color='black', linestyle=':', linewidth=1)
            
    else:
        res = []
        x_labels = []
        for table_name in sheet.tables:
            summo = 0
            table = sheet.tables[table_name]
            if sheet_name == "Encryptors":
                x_labels.append(table_name.split("_")[0])
                plt.xlabel('Number of encryptors')
            elif sheet_name == "Looping":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Loops')
            elif sheet_name == "message_Size":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Size duplication')
            elif sheet_name == "parallel_Opening":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Openings')
            elif sheet_name == "exclusive_Opening":
                x_labels.append(table_name.split("_")[1])
                plt.xlabel('Openings')
            elif sheet_name == "parallel_Opening_Closing":
                x_labels.append(table_name.split("_")[2])
                plt.xlabel('Openings and Closings')
            elif sheet_name == "exclusive_Opening_Closing":
                x_labels.append(table_name.split("_")[2])
                plt.xlabel('Openings and Closings')
            header_values = [cell.value for cell in sheet[table.ref][0]]
            gasses = header_values.index("Gas")
            for row in sheet[table.ref][1:]:
                summo += row[gasses].value
            res.append(int(summo))
        
        plt.plot(x_labels, res, color=(random.random(), random.random(), random.random()), label='Healthcare process')
        
        
        
    y_ticks = np.linspace(min(res), max(res), 7)
    plt.yticks(y_ticks)
    plt.gca().yaxis.set_major_formatter(FuncFormatter(lambda x, _: f'{x:,.0f}'.replace(',', '.')))

    min_val = min(res)
    plt.fill_between(range(len(res)), min_val, res, color=(random.random(), random.random(), random.random()), alpha=0.3) 
    plt.grid(axis='y', linestyle='-', alpha=0.5)
    plt.ylabel('Cumulative gas used')
    plt.legend()
    plt.show()
    




